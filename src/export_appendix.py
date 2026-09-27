"""Export Appendix 2 (preliminary findings) to Excel: output/Appendix2_PERT_preliminary.xlsx.

Every number is read from the analysis outputs (run replica.py --core proxy, val_attrib.py,
style_check.py, macro_v2.py and preqin.py first); nothing is typed by hand. Inputs are blue,
derived cells are Excel formulas (black), so differences, capture ratios and totals recompute
if an input is edited. The file is set to recalculate fully when opened.
"""
import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import macrobond_pull as mb
import style_check as sc
from data import ROOT
from macro_v2 import LABEL, drivers, targets

OUT = ROOT / "output" / "Appendix2_PERT_preliminary.xlsx"
O = ROOT / "output"

F = "Arial"
BLUE, BLACK, GREY = "0000FF", "000000", "595959"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")


def font(bold=False, color=BLACK, size=10, italic=False):
    return Font(name=F, bold=bold, color=color, size=size, italic=italic)


class Sheet:
    """Small helper: writes tables top to bottom and remembers where things are."""

    def __init__(self, wb, name, title, widths):
        self.ws = wb.create_sheet(name)
        self.row = 1
        self.ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.text(title, bold=True, size=13)
        self.row += 1

    def text(self, s, bold=False, size=10, italic=False, color=BLACK):
        c = self.ws.cell(self.row, 1, s)
        c.font = font(bold, color, size, italic)
        self.row += 1

    def header(self, cols):
        for j, h in enumerate(cols, 1):
            c = self.ws.cell(self.row, j, h)
            c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
            c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal="center" if j > 1 else "left", vertical="center", wrap_text=True)
        self.ws.row_dimensions[self.row].height = 30
        self.row += 1

    def line(self, values, fmts=None, bold=False, sub=False):
        """values: label, then numbers or formula strings ('=...'). Numbers are blue inputs."""
        r = self.row
        for j, v in enumerate(values, 1):
            if v is None or (isinstance(v, float) and np.isnan(v)):
                v = None
            c = self.ws.cell(r, j, v)
            is_formula = isinstance(v, str) and v.startswith("=")
            color = BLACK if (j == 1 or is_formula or isinstance(v, str)) else BLUE
            c.font = font(bold, color)
            c.border = Border(bottom=THIN)
            if sub:
                c.fill = SUB_FILL
            if j > 1:
                c.alignment = Alignment(horizontal="right")
                if fmts and fmts[j - 2]:
                    c.number_format = fmts[j - 2]
        self.row += 1
        return r

    def gap(self, n=1):
        self.row += n


PCT = '0.0%;-0.0%;"-"'
PP = '+0.0;-0.0;0.0'
X2 = "0.00"


def notes(wb):
    s = Sheet(wb, "Notes", "Appendix 2: Preliminary findings, data (MSCI PERT, Proxy PERT, World IMI)", [110])
    for t in [
        "Terms",
        "MSCI PERT: MSCI World Private Equity Return Tracker Index (the index GTPE tracks), built by MSCI from listed stocks.",
        "Proxy PERT: our rebuild of MSCI PERT from published MSCI indexes, Cambridge Associates PE/VC sector weights and MSCI's Aug 2025 PE sector snapshot.",
        "World IMI: MSCI World Investable Market Index (large, mid and small caps; net returns, USD).",
        "",
        "Conventions",
        "Blue numbers are results read from the analysis; black cells are formulas. Returns in % per year unless stated.",
        "Attribution tables use sums of monthly returns (additive); return tables use compounded annual returns.",
        "",
        "Sources",
        "Bloomberg (MSCI indexes, PERT components, valuation data, US yields/spreads/dollar); Macrobond (SPF, OECD leading indicator, ISM, Chicago Fed NFCI, 5y5y breakeven, Brent);",
        "Preqin Pro exports (closed-end fund indexes, quarterly TWR, USD); Cambridge Associates benchmark commentaries; MSCI 'Rethinking Access to Private Equity' (Oct 2025);",
        "Kenneth French data library (Fama-French factors); AQR data sets (Quality Minus Junk).",
        "",
        "Caveats",
        "Few macro cycles since 2007; sensitivities are co-movements, not causal estimates. MSCI PERT history before 2025 is back-tested.",
        "Preqin indexes are net of fees and appraisal-based; the listed trackers are gross, market-priced returns.",
        "Licensed data (MSCI, Bloomberg, Preqin, Macrobond): internal use only.",
    ]:
        s.text(t, bold=t in ("Terms", "Conventions", "Sources", "Caveats"))


def replication(wb):
    s = Sheet(wb, "Replication", "2.1 Replication: returns on MSCI's published window (29 Dec 2006 - 31 Mar 2025)", [46, 10, 10, 10, 10, 10, 12])
    rec = sc.reconcile()
    s.header(["", "1Y", "3Y", "5Y", "10Y", "Full", "Risk (quarterly)"])
    names = {"MSCI PERT": "MSCI PERT (our data; matches MSCI published)",
             "Proxy PERT (fixed weights)": "Proxy PERT, fixed weights",
             "Proxy PERT (rebalanced, from Dec 2008)": "Proxy PERT, rebalanced quarterly (from Dec 2008)",
             "World IMI (net)": "World IMI (our data; matches MSCI published)"}
    for k, lab in names.items():
        r = rec.loc[k]
        full = r["Full"] / 100 if not k.startswith("Proxy PERT (rebal") else None
        risk = r["Risk (quarterly)"] / 100 if not k.startswith("Proxy PERT (rebal") else None
        s.line([lab, r["1Y"] / 100, r["3Y"] / 100, r["5Y"] / 100, r["10Y"] / 100, full, risk], [PCT] * 6)
    s.gap()
    s.text("MSCI published (Rethinking Access to Private Equity, p.14): World PERT 4.1 / 10.4 / 16.9 / 10.4 / 9.9, risk 19.1;"
           " World IMI 6.3 / 6.9 / 15.8 / 9.2 / 7.0, risk 17.6.", italic=True, color=GREY)
    s.gap()
    w = pd.read_csv(O / "replica_proxy" / "world_monthly.csv", index_col=0, parse_dates=True).loc["2008-12-31":]
    d = w[["MSCI PERT", "Proxy PERT dynamic"]].dropna()
    s.header(["Proxy PERT (rebalanced) vs MSCI PERT, Dec 2008 - Aug 2026", "Value"])
    s.line(["Correlation (monthly)", d.corr().iloc[0, 1]], [X2])
    s.line(["Tracking error (annualised)", (d.iloc[:, 0] - d.iloc[:, 1]).std() * np.sqrt(12)], [PCT])


def drivers_sheet(wb):
    s = Sheet(wb, "Return drivers", "2.2 What drove MSCI PERT's lead over World IMI (Dec 2008 - Aug 2026, % per year)", [34, 9, 10, 12, 10, 10, 10, 9, 12])
    v = pd.read_csv(O / "val_attrib" / "world.csv", index_col=[0, 1])
    full = v.xs("Dec 2008 - Aug 2026, per year", level=0)
    cols = ["Total", "Dividends", "Earnings growth", "P/E change", "Currency", "Style tilts", "Other", "Unexplained"]

    def parts(row):
        tilts = row[["value", "growth", "momentum", "minvol", "leverage", "size"]].sum()
        return [row["income"], row["earnings"], row["multiple"] + row["unsplit"], row["currency"], tilts, row["imi_gap"], row["unexplained"]]

    s.header([""] + cols)
    rows = {}
    for name in ["MSCI PERT", "World IMI"]:
        p = parts(full.loc[name])
        r = s.row
        rows[name] = r
        s.line([name, f"=SUM(C{r}:I{r})"] + [x / 100 for x in p], [PCT] * 8)
    a, b = rows["MSCI PERT"], rows["World IMI"]
    s.line(["Difference (MSCI PERT - World IMI)"] + [f"={get_column_letter(j)}{a}-{get_column_letter(j)}{b}" for j in range(2, 10)],
           [PCT] * 8, bold=True, sub=True)
    s.text("Other = small-cap gap (regional IMI vs large/mid-cap index). Unexplained = MSCI PERT minus Proxy PERT.", italic=True, color=GREY)
    s.gap()

    br = pd.read_csv(O / "replica_proxy" / "bridge_pert_vs_imi.csv", index_col=0).iloc[:, 0]
    s.header(["Same lead, by exposure (% per year)", "Contribution"])
    first = s.row
    for lab, key in [("Sector mix (PE-matched Core minus regional market)", "sector mix (PE-matched Core minus regional IMI)"),
                     ("Regional mix", "regional mix (PERT's regional weights vs World IMI's)"),
                     ("Market exposure above 1", "market exposure above/below 1"),
                     ("Style tilts", "style tilts"),
                     ("Unexplained (MSCI PERT minus Proxy PERT)", "unexplained (PERT minus Proxy PERT)")]:
        s.line([lab, br[key] / 100], [PCT])
    s.line(["Total", f"=SUM(B{first}:B{s.row - 1})"], [PCT], bold=True, sub=True)
    s.gap()

    s.header(["Difference by period (% per year)", "Total", "Dividends", "Earnings growth", "P/E change", "Currency", "Style tilts", "Other", "Unexplained"])
    for per in ["2009-2014", "2015-2021", "2022-Aug 2026"]:
        p = parts(v.loc[(per, "Difference")])
        r = s.row
        s.line([per, f"=SUM(C{r}:I{r})"] + [x / 100 for x in p], [PCT] * 8)
    s.gap()

    ff = pd.read_csv(O / "style_check" / "fama_french.csv", index_col=[0, 1]).loc[("US", "2007-2026")]
    qmj = sc.read_aqr("QMJ").loc["2007-01":].mean() * 12
    s.header(["Published US factor returns, 2007-2026 (long-short, % per year)", "Return"])
    for lab, val in [("Value (Fama-French HML)", ff["HML"] / 100), ("Size, small minus big (Fama-French SMB)", ff["SMB"] / 100),
                     ("Profitability (Fama-French RMW)", ff["RMW"] / 100), ("Momentum (Fama-French)", ff["Mom"] / 100),
                     ("Quality minus junk (AQR QMJ)", qmj)]:
        s.line([lab, val], [PCT])
    s.text("PE tilts towards value, small size and leverage; these lagged over the period, consistent with the small style drag.", italic=True, color=GREY)


def macro_sheet(wb):
    s = Sheet(wb, "Macro", "2.3 Macro sensitivities: MSCI PERT vs World IMI (monthly, Jan 2007 - Aug 2026)", [52, 18, 11, 8, 11, 8, 12, 14])
    u = pd.read_csv(O / "macro_v2" / "univariate.csv")
    D = drivers().reindex(targets().loc["2007-01":].index)
    sd = D.std()
    units = {"hy_spread": ("bp", 100), "real_yield": ("bp", 100), "be_5y5y": ("bp", 100), "ust_2y": ("bp", 100),
             "dollar": ("%", 1), "oil": ("%", 1), "fwd_eps_rev": ("%", 1), "nfci": ("index pts", 1), "ism_new_orders": ("pts", 1),
             "cli_g7": ("index pts", 1)}
    order = [("hy_spread", "US high-yield spread (OAS)"),
             ("nfci", "US financial conditions (Chicago Fed NFCI; higher = tighter)"),
             ("dollar", "US dollar index"),
             ("real_yield", "US 10y real yield (TIPS)"),
             ("ism_new_orders", "ISM manufacturing new orders"),
             ("be_5y5y", "5y5y inflation expectations (breakeven)"),
             ("fwd_eps_rev", "MSCI World 12m-forward EPS revision (analysts)"),
             ("oil", "Brent crude"),
             ("ust_2y", "US 2y Treasury yield (policy path)")]
    s.header(["Driver", "Typical monthly move (1 s.d.)", "MSCI PERT", "t-stat", "World IMI", "t-stat", "PERT minus IMI"])
    for key, lab in order:
        unit, mult = units[key]
        g = u[u["driver"] == LABEL[key]].set_index("series")
        r = s.row
        move = f"+{sd[key] * mult:.2f} {unit}" if unit != "bp" else f"+{sd[key] * mult:.0f} bp"
        s.line([lab, move, g.loc["MSCI PERT", "beta"] / 100, g.loc["MSCI PERT", "t"], g.loc["World IMI", "beta"] / 100,
                g.loc["World IMI", "t"], f"=C{r}-E{r}"], [None, PCT, "0.0", PCT, "0.0", PCT])
    s.text("Each driver on its own; return in the same month for a typical (one standard deviation) move. t-stats are Newey-West.", italic=True, color=GREY)
    s.gap()

    m = pd.read_csv(O / "macro_v2" / "multivariate.csv", index_col=0)
    def split(x):  # "+0.18 (+0.7)" -> (0.0018, 0.7)
        b, t = str(x).replace("(", "").replace(")", "").split()
        return float(b) / 100, float(t)
    s.header(["All drivers together (return per typical move)", "MSCI PERT", "t-stat", "World IMI", "t-stat"])
    for col in [c for c in m.columns if c not in ("R2", "n")]:
        (bp, tp), (bi, ti) = split(m.loc["MSCI PERT", col]), split(m.loc["World IMI", col])
        s.line([col, bp, tp, bi, ti], [PCT, "0.0", PCT, "0.0"])
    s.line(["R-squared", float(m.loc["MSCI PERT", "R2"]), None, float(m.loc["World IMI", "R2"]), None], [X2, None, X2, None])
    s.text("With all drivers together, only credit spreads, the dollar and real yields remain significant.", italic=True, color=GREY)
    s.gap()

    # growth expectations (quarterly SPF), per +0.25pp revision
    T = targets()
    q = (1 + T).resample("QE").prod(min_count=3) - 1
    mq = mb.quarterly()
    X = pd.concat([mq[["spf_gdp_q1", "spf_gdp_q2", "spf_gdp_q3", "spf_gdp_q4"]].mean(axis=1).diff().rename("g"),
                   mq["spf_cpi_next_year"].diff().rename("c")], axis=1)
    s.header(["Forecasters' growth expectations (SPF, quarterly 2007-2026)", "Per +1pp revision", "t-stat", "Per +0.25pp revision"])
    for name in ["MSCI PERT", "World IMI"]:
        for excl in (False, True):
            z = pd.concat([q[name].rename("y"), X], axis=1).loc["2007":].dropna()
            if excl:
                z = z[z.index.year != 2020]
            mdl = sm.OLS(z["y"], sm.add_constant(z[["g", "c"]])).fit(cov_type="HAC", cov_kwds={"maxlags": 2})
            r = s.row
            s.line([f"{name}{' (excluding 2020)' if excl else ''}", mdl.params["g"], mdl.tvalues["g"], f"=B{r}*0.25"], [PCT, "0.0", PCT])
    s.text("Revision = change in the SPF mean expected real GDP growth over the next four quarters; controls for the revision to next-year CPI expectations.",
           italic=True, color=GREY)
    s.text("Chicago Fed NFCI: weekly composite of ~105 US money-market, debt, equity and banking indicators; 0 = average since 1971, positive = tighter than average.",
           italic=True, color=GREY)


def preqin_sheet(wb):
    s = Sheet(wb, "Reported PE (Preqin)", "2.4 Reported PE (Preqin) vs listed trackers (Q1 2007 - Q1 2026)", [44, 12, 16, 11, 13, 15])
    lv = pd.read_csv(O / "preqin" / "levels.csv")
    pick = [("Preqin Global Private Equity & Venture Capital", "MSCI PERT", "Global PE & VC vs MSCI PERT / World IMI"),
            ("Preqin Global Private Equity & Venture Capital", "Proxy PERT", "Global PE & VC vs Proxy PERT / World IMI"),
            ("Preqin North America Private Equity", "NA trackers", "North America PE vs MSCI NA trackers / NA IMI"),
            ("Preqin Europe Private Equity", "Europe trackers", "Europe PE vs MSCI Europe trackers / Europe IMI"),
            ("Preqin Asia Private Equity", "Pacific trackers", "Asia PE (incl. emerging Asia) vs MSCI Pacific trackers / Pacific IMI")]
    s.header(["Annual return", "Reported PE", "Tracker", "Public market", "Tracker's share of PE's lead"])
    for pe, tr, lab in pick:
        row = lv[(lv["PE index"] == pe) & (lv["tracker"] == tr)].iloc[0]
        r = s.row
        s.line([lab, row["PE %/yr"] / 100, row["tracker %/yr"] / 100, row["market %/yr"] / 100, f"=(C{r}-D{r})/(B{r}-D{r})"], [PCT] * 4)
    s.text("Tracker blends use PERT's average component weights (North America 62/38, Europe 85/15 buyout/VC; Pacific 50/50).", italic=True, color=GREY)
    s.gap()

    g = lv[(lv["PE index"] == pick[0][0]) & (lv["tracker"] == "MSCI PERT")].iloc[0]
    s.header(["Reported global PE minus MSCI PERT (% per year)", "Gap"])
    for lab, col in [("2007-2014", "PE - tracker 07-14"), ("2015-2021", "PE - tracker 15-21"), ("2022 - Q1 2026", "PE - tracker 22-26")]:
        s.line([lab, g[col] / 100], [PCT])
    s.gap()

    sm_ = pd.read_csv(O / "preqin" / "smoothing.csv")
    lg = pd.read_csv(O / "preqin" / "lags.csv")
    a = sm_[sm_["PE index"] == pick[0][0]].iloc[0]
    l = lg[(lg["PE index"] == pick[0][0]) & (lg["tracker"] == "MSCI PERT")].iloc[0]
    s.header(["Smoothing and lag (Global PE & VC vs MSCI PERT)", "Value"])
    s.line(["Volatility, as reported", a["vol PE, reported %"] / 100], [PCT])
    s.line(["Volatility, desmoothed (Geltner)", a["vol PE, desmoothed %"] / 100], [PCT])
    s.line(["Volatility, MSCI PERT", a["vol tracker %"] / 100], [PCT])
    s.line(["Max drawdown, reported PE", a["max DD PE %"] / 100], [PCT])
    s.line(["Max drawdown, MSCI PERT (quarterly)", a["max DD tracker %"] / 100], [PCT])
    first = s.row
    for k, lab in enumerate(["same quarter", "one quarter later", "two quarters later", "three quarters later"]):
        s.line([f"Share of a tracker move showing up in reported PE: {lab}", l[f"lag {k}"]], [PCT])
    s.line(["Share within a year", f"=SUM(B{first}:B{s.row - 1})"], [PCT], bold=True, sub=True)
    s.gap()

    nc = pd.read_csv(O / "preqin" / "nowcast.csv")
    n = nc[nc["PE index"] == pick[0][0]].set_index("tracker")
    s.header(["Nowcast of reported global PE, out of sample 2012 - Q1 2026", "R-squared", "Error (RMSE)"])
    s.line(["Using MSCI PERT (current + last two quarters)", n.loc["MSCI PERT", "OOS R2 vs hist. mean (tracker)"], n.loc["MSCI PERT", "RMSE % (tracker)"] / 100], [X2, PCT])
    s.line(["Using World IMI", n.loc["MSCI PERT", "OOS R2 (market instead)"], n.loc["MSCI PERT", "RMSE % (market instead)"] / 100], [X2, PCT])
    s.line(["Using Proxy PERT", n.loc["Proxy PERT", "OOS R2 vs hist. mean (tracker)"], n.loc["Proxy PERT", "RMSE % (tracker)"] / 100], [X2, PCT])
    s.line(["Naive: historical average", 0.0, n.loc["MSCI PERT", "RMSE, hist. mean % (tracker)"] / 100], [X2, PCT])
    s.text("Each quarter is predicted with a regression fitted only on earlier quarters. R-squared is relative to the historical-average forecast.",
           italic=True, color=GREY)


def main():
    wb = Workbook()
    wb.remove(wb.active)
    notes(wb)
    replication(wb)
    drivers_sheet(wb)
    macro_sheet(wb)
    preqin_sheet(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
